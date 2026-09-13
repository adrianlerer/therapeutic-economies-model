#!/usr/bin/env python3
"""Build a publication-style DOCX from the canonical Markdown manuscript."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "paper/MANUSCRIPT.md"
OUTPUT = ROOT / "paper/Therapeutic_Economies_Preprint.docx"
FIGURE = ROOT / "results/trajectories.png"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, value, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(value.strip())
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0, 0, 0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def strip_markup(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"`(.*?)`", r"\1", text)
    replacements = {
        r"\operatorname{clip}": "clip",
        r"\tilde{x}": "x̃",
        r"\frac": "",
        r"\exp": "exp",
        r"\alpha": "α", r"\beta": "β", r"\gamma": "γ",
        r"\delta": "δ", r"\varepsilon": "ε", r"\rho": "ρ",
        r"\eta": "η", r"\theta": "θ", r"\phi": "φ",
        r"\ell": "ℓ", r"\kappa": "κ", r"\lambda": "λ",
        r"\nu": "ν", r"\mu": "μ", r"\pi": "π", r"\geq": "≥",
        r"\leq": "≤", r"\in": "∈", r"\{": "(", r"\}": ")",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = text.replace("\\(", "").replace("\\)", "")
    text = text.replace("_{t+1}", "ₜ₊₁").replace("_t", "ₜ").replace("_T", "T")
    text = text.replace("_C", "C").replace("_P", "P").replace("_i", "i").replace("_0", "₀")
    text = text.replace("^", "^").replace("\\", "")
    text = text.replace("{", "(").replace("}", ")")
    return text


def configure(document):
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, before, after in (
        ("Title", 20, 0, 6),
        ("Subtitle", 13, 0, 8),
        ("Heading 1", 14, 12, 5),
        ("Heading 2", 11.5, 9, 4),
    ):
        style = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Aptos Display" if name in ("Title", "Heading 1") else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        ppr = style.element.get_or_add_pPr()
        border = ppr.find(qn("w:pBdr"))
        if border is not None:
            ppr.remove(border)

    header = section.header.paragraphs[0]
    header.text = "THERAPEUTIC ECONOMIES  |  PREPRINT"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(80, 80, 80)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(footer)


def add_rich_paragraph(document, text, style=None, abstract=False):
    p = document.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if abstract:
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.right_indent = Inches(0.28)
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        italic = part.startswith("*") and part.endswith("*") and not bold
        code = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if bold else part[1:-1] if italic or code else part
        clean = strip_markup(clean)
        r = p.add_run(clean)
        r.bold = bold
        r.italic = italic
        r.font.name = "Aptos Mono" if code else "Aptos"
        r.font.size = Pt(10 if abstract else 10.5)
        r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc)
    in_abstract = False
    i = 0
    title_seen = False
    subtitle_seen = False
    figure_inserted = False

    while i < len(lines):
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = [[strip_markup(c.strip()) for c in item.strip("|").split("|")] for item in table_lines]
            rows = [rows[0]] + rows[2:]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, value in enumerate(row):
                    set_cell_text(table.cell(ri, ci), value, bold=ri == 0)
                    if ri == 0:
                        shade(table.cell(ri, ci), "D9D9D9")
            if not figure_inserted and any("Inversion frequency" in c for c in rows[0]):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(FIGURE), width=Inches(6.3))
                cap = doc.add_paragraph("Figure 1. Mean endogenous-capability trajectories across 200 seeds.")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True
                cap.runs[0].font.size = Pt(9)
                figure_inserted = True
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run(strip_markup(line[2:]))
            title_seen = True
        elif line.startswith("## "):
            heading = strip_markup(line[3:])
            if title_seen and not subtitle_seen:
                p = doc.add_paragraph(style="Subtitle")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(heading)
                subtitle_seen = True
            else:
                p = doc.add_paragraph(heading, style="Heading 1")
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            in_abstract = heading == "Abstract"
        elif line.startswith("### "):
            p = doc.add_paragraph(strip_markup(line[4:]), style="Heading 2")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif line.startswith("- "):
            p = add_rich_paragraph(doc, line[2:], style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif re.match(r"^\d+\. ", line):
            p = add_rich_paragraph(doc, re.sub(r"^\d+\. ", "", line), style="List Number")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif line.startswith("\\["):
            formula = [line.replace("\\[", "")]
            while i + 1 < len(lines):
                i += 1
                next_line = lines[i].strip()
                if next_line == "\\]":
                    break
                formula.append(next_line)
            p = doc.add_paragraph(strip_markup(" ".join(formula)))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Cambria Math"
                r.font.size = Pt(10)
        elif not title_seen:
            add_rich_paragraph(doc, line)
        elif in_abstract:
            p = add_rich_paragraph(doc, line, abstract=True)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif title_seen and subtitle_seen and len(doc.paragraphs) < 9 and not line.startswith("##"):
            p = add_rich_paragraph(doc, line)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = add_rich_paragraph(doc, line)
            if line.startswith("**Keywords:**"):
                for r in p.runs:
                    r.font.size = Pt(10)
        i += 1

    props = doc.core_properties
    props.title = "Therapeutic Economies: Protective Selection, Endogenous Capabilities, and Institutional Dependence"
    props.author = "Ignacio Adrián Lerer"
    props.subject = "Theoretical-computational preprint"
    props.keywords = "evolutionary game theory, capability formation, institutional dependence"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
